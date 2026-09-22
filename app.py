import io
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import pdfplumber
import streamlit as st

st.set_page_config(
    page_title="SPX Laporan Scan Converter", page_icon="📦", layout="wide"
)

st.title("📦 SPX Laporan Scan PDF ➡️ Excel Converter")

MARKING_MAP = {
    "Abepura DC": "DJJ-C1-1",
    "Alak DC": "KOE-C1-1",
    "Bacan Hub": "LAH-C1-1",
    "Baguala DC": "AMQ-C1-1",
    "Balikpapan DC": "BPN-C1-1",
    "Banjarmasin DC": "BDJ1-C1-1",
    "Banjarmasin 2 DC": "BDJ2-C1-1",
    "Banjarbaru DC": "BJB-C1-1",
    "Batam DC": "BTH-C1-1",
    "Dungingi DC": "GTO-C1-1",
    "Kalawat DC": "MDU-C1-1",
    "Kota Waingapu Hub": "WGP-C1-1",
    "Kota Waingapu 2 Hub": "WGP2-C1-1",
    "Kota Waingapu 4 Hub": "WGP4-C1-1",
    "Labuhan Bajo DC": "LBJ-C1-1",
    "Loli Hub": "TMC2-C1-1",
    "Loura (Laura) Hub": "TMC-C1-1",
    "Manokwari Barat DC": "MKW-C1-1",
    "Mantikulore DC": "PLW-C1-1",
    "Medan DC": "KNO-C1-1",
    "Medan Amplas DC": "KNO2-C1-1",
    "Medan Deli DC": "KNO3-C1-1",
    "Merauke DC": "MKQ-C1-1",
    "Mimika Baru Hub": "TIM-C1-1",
    "Nabire Hub": "NBX-C1-1",
    "Palangka Raya DC": "PKY-C1-1",
    "Percut Sei Tuan DC": "PST-C1-1",
    "Pekanbaru DC": "PKU-C1-1",
    "Pekanbaru 2 DC": "PKU2-C1-1",
    "Pontianak DC": "PNK-C1-1",
    "Pontianak 2 DC": "PNK2-C1-1",
    "Sungai Kakap DC": "PNK3-C1-1",
    "Sorong Utara DC": "SOQ-C1-1",
    "Tarakan Barat Hub": "TRKB-C1-1",
    "Tarakan Barat 4 Hub": "TRKB4-C1-1",
    "Tarakan Timur Hub": "TRKT-C1-1",
    "Tarakan Utara Hub": "TRKU-C1-1",
    "Teluk Mutiara Hub": "ARD-C1-1",
    "Ternate Hub": "TTE-C1-1",
    "Ternate Utara Hub": "TTU-C1-1",
    "Ternate Selatan Hub": "TTS-C1-1",
    "Ternate Selatan 2 Hub": "TTS2-C1-1",
    "Ternate Selatan 3 Hub": "TTS3-C1-1",
    "Wamena Hub": "WMX-C1-1",
    "Wua-Wua DC": "KDI-C1-1",
}


def apply_table_formatting(ws, start_row, max_col):
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    for col in range(1, max_col + 1):
        cell = ws.cell(row=start_row, column=col)
        cell.font = Font(bold=True, name="Calibri")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for r in range(start_row, ws.max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            if r > start_row:
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        if col[0].column > max_col:
            continue
        for cell in col:
            if cell.row < start_row:
                continue
            if cell.value is not None:
                val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)


def extract_data_from_pdf(pdf_file):
    extracted_rows = []
    valid_destinations = list(MARKING_MAP.keys())

    with pdfplumber.open(pdf_file) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue

            lt_match = re.search(r"\b(LT[A-Z0-9]{8,})\b", text)
            lt_num = lt_match.group(1) if lt_match else ""

            dest = ""
            for valid_dest in valid_destinations:
                if valid_dest in text and valid_dest != "SURABAYA DC":
                    dest = valid_dest
                    break

            if not dest:
                all_dcs = re.findall(
                    r"\b([A-Za-z0-9\s-]+?\s*(?:DC|Hub))\b", text, re.IGNORECASE
                )
                for d in all_dcs:
                    d_clean = d.strip()
                    if "SURABAYA" not in d_clean.upper():
                        dest = d_clean
                        break

            std_match = re.search(r"(\d{4}/\d{2}/\d{2})\s*\d{2}:\d{2}:\d{2}STD", text)
            if not std_match:
                std_match = re.search(r":\s*(\d{4}/\d{2}/\d{2})", text)
            tgl = (
                std_match.group(1).replace("/", "-")
                if std_match
                else "2026-09-14"
            )

            lines = text.split("\n")
            for line in lines:
                to_match = re.search(r"\b(TO\d{8}[A-Z0-9]+)\b", line)
                if to_match:
                    to_num = to_match.group(1)

                    clean_line = re.sub(r"\d{4}/\d{2}/\d{2}", "", line)
                    clean_line = re.sub(r"\d{2}:\d{2}:\d{2}", "", clean_line)

                    weights = re.findall(r"\b(\d{1,3}[\.,]\d{1,3})\b", clean_line)

                    gw = 0.0
                    if weights:
                        try:
                            val_str = weights[0].replace(",", ".")
                            gw = round(float(val_str), 3)
                        except ValueError:
                            gw = 0.0

                    extracted_rows.append({
                        "TGL": tgl,
                        "Vendor": "Lion Parcel",
                        "Sc Origin": "SURABAYA DC",
                        "Sc Destination": dest,
                        "Lt Number": lt_num,
                        "To Number": to_num,
                        "Gross Weight": gw,
                        "Remarks": "BAG",
                    })

    df_extracted = pd.DataFrame(extracted_rows)

    if not df_extracted.empty:
        df_extracted = df_extracted.iloc[::-1].reset_index(drop=True)

        marking_list = []
        dest_counters = {}

        for _, row in df_extracted.iterrows():
            dest_name = row["Sc Destination"]
            base_marking = MARKING_MAP.get(dest_name, "C1-1")

            count = dest_counters.get(dest_name, 0)
            batch_num = (count // 15) + 1
            dest_counters[dest_name] = count + 1

            parts = base_marking.rsplit("-", 1)
            if len(parts) == 2 and parts[1].isdigit():
                new_marking = f"{parts[0]}-{batch_num}"
            else:
                new_marking = f"{base_marking}-{batch_num}"

            marking_list.append(new_marking)

        df_extracted["Marking"] = marking_list

    return df_extracted


uploaded_file = st.file_uploader("Upload File PDF SPX", type=["pdf"])

if uploaded_file is not None:
    df = extract_data_from_pdf(uploaded_file)

    if not df.empty:
        total_rows = len(df)
        total_gw = round(df["Gross Weight"].sum(), 3)
        st.success(f"Berhasil! Total **{total_rows}** TO | Total GW: **{total_gw:,.3f}** kg")

        df_sjm = pd.DataFrame({
            "TGL": df["TGL"],
            "Vendor": "LION PARCEL",
            "SC Orgin": df["Sc Origin"],
            "DESTINATION": df["Sc Destination"],
            "LT NUMBER": df["Lt Number"],
            "TO NUMBER": df["To Number"],
            "Gross Weight": df["Gross Weight"],
            "REMAKE": df["Remarks"],
            "TOTAL": "",
        })

        df_marking = pd.DataFrame({
            "Tanggal": df["TGL"],
            "Vendor": df["Vendor"],
            "Sc Origin": df["Sc Origin"],
            "Sc Destination": df["Sc Destination"],
            "Lt Number": df["Lt Number"],
            "To Number": df["To Number"],
            "Marking": df["Marking"],
            "Gross Weight": df["Gross Weight"],
            "Remarks": df["Remarks"],
            "External Number": df["Marking"]
            + "/"
            + df["Lt Number"]
            + "/"
            + df["Remarks"],
            "Clear Gw": df["Gross Weight"],
        })

        # Aggregation PVT
        df_pvt = (
            df_marking.groupby("External Number", sort=False)
            .agg(
                Count_of_External_Number=("To Number", "count"),
                Sum_of_Clear_Gw=("Clear Gw", lambda x: round(x.sum(), 3)),
            )
            .reset_index()
        )
        df_pvt.columns = [
            "External Number",
            "Count of External Number",
            "Sum of Clear Gw",
        ]

        # Aggregation Sheet3
        df_sheet3 = (
            df.groupby("Sc Destination", sort=False)
            .agg(
                Count_of_To_Number=("To Number", "count"),
                Sum_of_Gross_Weight=("Gross Weight", lambda x: round(x.sum(), 3)),
            )
            .reset_index()
        )
        df_sheet3.columns = [
            "Sc Destination",
            "Count of To Number",
            "Sum of Gross Weight",
        ]

        t1, t2, t3, t4 = st.tabs(
            ["📋 Sheet SJM", "🏷️ Sheet MARKING", "📑 Sheet PVT", "📊 Sheet3"]
        )

        with t1:
            st.markdown(
                "**SURAT JALAN MANUAL SURABAYA DC VIA LION STD | 14 SEPTEMBER 2026 TRIP 2**"
            )
            st.dataframe(df_sjm, use_container_width=True, hide_index=True)

        with t2:
            st.markdown("**MARKING SPX OSO SUB DC CYCLE | 14 SEPTEMBER 2026 TRIP 2**")
            st.dataframe(df_marking, use_container_width=True, hide_index=True)

        with t3:
            st.dataframe(df_pvt, use_container_width=True, hide_index=True)

        with t4:
            st.dataframe(df_sheet3, use_container_width=True, hide_index=True)

        wb = openpyxl.Workbook()
        red_fill = PatternFill(
            start_color="FF0000", end_color="FF0000", fill_type="solid"
        )
        yellow_fill = PatternFill(
            start_color="FFFF00", end_color="FFFF00", fill_type="solid"
        )
        grey_fill = PatternFill(
            start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
        )
        thin_border = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000"),
        )

        # 1. SHEET SJM
        ws_sjm = wb.active
        ws_sjm.title = "SJM"
        ws_sjm.merge_cells("A1:H1")
        cell_r1 = ws_sjm.cell(
            row=1, column=1, value="SURAT JALAN MANUAL SURABAYA DC VIA LION STD"
        )
        cell_r1.font = Font(bold=True, size=11, name="Calibri")
        cell_r1.alignment = Alignment(horizontal="center", vertical="center")

        ws_sjm.merge_cells("A2:H2")
        cell_r2 = ws_sjm.cell(row=2, column=1, value="14 SEPTEMBER 2026 TRIP 2")
        cell_r2.font = Font(bold=True, size=11, name="Calibri")
        cell_r2.alignment = Alignment(horizontal="center", vertical="center")

        total_cell = ws_sjm.cell(row=2, column=9, value=total_rows)
        total_cell.font = Font(bold=True, size=11, name="Calibri")
        total_cell.alignment = Alignment(horizontal="center", vertical="center")

        for c_idx, col_name in enumerate(df_sjm.columns, 1):
            ws_sjm.cell(row=3, column=c_idx, value=col_name)

        for r_idx, row_val in enumerate(df_sjm.itertuples(index=False), 4):
            for c_idx, val in enumerate(row_val, 1):
                ws_sjm.cell(row=r_idx, column=c_idx, value=val)

        sjm_last_row = total_rows + 3
        sjm_gt_row = sjm_last_row + 1
        ws_sjm.merge_cells(
            start_row=sjm_gt_row, start_column=1, end_row=sjm_gt_row, end_column=6
        )
        gt_sjm_label = ws_sjm.cell(row=sjm_gt_row, column=1, value="Grand Total")
        gt_sjm_label.font = Font(bold=True, name="Calibri")
        gt_sjm_label.alignment = Alignment(
            horizontal="center", vertical="center"
        )

        gt_sjm_val = ws_sjm.cell(
            row=sjm_gt_row, column=7, value=f"=SUM(G4:G{sjm_last_row})"
        )
        gt_sjm_val.font = Font(bold=True, name="Calibri")
        gt_sjm_val.alignment = Alignment(
            horizontal="center", vertical="center"
        )

        apply_table_formatting(ws_sjm, start_row=3, max_col=len(df_sjm.columns))
        for r in range(1, 3):
            for c in range(1, 10):
                ws_sjm.cell(row=r, column=c).border = thin_border

        # 2. SHEET MARKING
        ws_mk = wb.create_sheet(title="MARKING")
        ws_mk.merge_cells("A1:K1")
        cell_mk1 = ws_mk.cell(
            row=1, column=1, value="MARKING SPX OSO SUB DC CYCLE"
        )
        cell_mk1.font = Font(bold=True, color="FFFF00", name="Calibri", size=11)
        cell_mk1.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(1, 12):
            cell = ws_mk.cell(row=1, column=c)
            cell.fill = red_fill
            cell.border = thin_border

        ws_mk.merge_cells("A2:K2")
        cell_mk2 = ws_mk.cell(row=2, column=1, value="14 SEPTEMBER 2026 TRIP 2")
        cell_mk2.font = Font(bold=True, color="FF0000", name="Calibri", size=10)
        cell_mk2.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(1, 12):
            cell = ws_mk.cell(row=2, column=c)
            cell.fill = yellow_fill
            cell.border = thin_border

        for c_idx, col_name in enumerate(df_marking.columns, 1):
            ws_mk.cell(row=3, column=c_idx, value=col_name)

        for r_idx, r in enumerate(df_marking.itertuples(index=False), 4):
            for c_idx, val in enumerate(r, 1):
                ws_mk.cell(row=r_idx, column=c_idx, value=val)

        mk_last_row = total_rows + 3
        footer_row = mk_last_row + 1
        ws_mk.merge_cells(
            start_row=footer_row, start_column=1, end_row=footer_row, end_column=7
        )
        gt_mk_label = ws_mk.cell(row=footer_row, column=1, value="Grand Total")
        gt_mk_label.font = Font(bold=True, color="FFFFFF", name="Calibri")
        gt_mk_label.alignment = Alignment(horizontal="center", vertical="center")

        gt_mk_val1 = ws_mk.cell(
            row=footer_row, column=8, value=f"=SUM(H4:H{mk_last_row})"
        )
        gt_mk_val1.font = Font(bold=True, color="FFFFFF", name="Calibri")
        gt_mk_val1.alignment = Alignment(horizontal="center", vertical="center")

        gt_mk_val2 = ws_mk.cell(
            row=footer_row, column=11, value=f"=SUM(K4:K{mk_last_row})"
        )
        gt_mk_val2.font = Font(bold=True, color="FFFFFF", name="Calibri")
        gt_mk_val2.alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, 12):
            cell = ws_mk.cell(row=footer_row, column=c)
            cell.fill = red_fill
            cell.border = thin_border

        apply_table_formatting(
            ws_mk, start_row=3, max_col=len(df_marking.columns)
        )

        # 3. SHEET PVT
        ws_pvt = wb.create_sheet(title="PVT")
        for c_idx, col_name in enumerate(df_pvt.columns, 1):
            cell = ws_pvt.cell(row=1, column=c_idx, value=col_name)
            cell.fill = grey_fill

        for p_idx, r in enumerate(df_pvt.itertuples(index=False), 2):
            ws_pvt.cell(row=p_idx, column=1, value=r[0])
            ws_pvt.cell(row=p_idx, column=2, value=r[1])
            ws_pvt.cell(row=p_idx, column=3, value=r[2])

        last_pvt_row = len(df_pvt) + 1
        gt_row = last_pvt_row + 1

        gt_cell1 = ws_pvt.cell(row=gt_row, column=1, value="Grand Total")
        gt_cell2 = ws_pvt.cell(
            row=gt_row, column=2, value=f"=SUM(B2:B{last_pvt_row})"
        )
        gt_cell3 = ws_pvt.cell(
            row=gt_row, column=3, value=f"=SUM(C2:C{last_pvt_row})"
        )

        for cell in (gt_cell1, gt_cell2, gt_cell3):
            cell.font = Font(bold=True, name="Calibri")
            cell.fill = grey_fill

        apply_table_formatting(
            ws_pvt, start_row=1, max_col=len(df_pvt.columns)
        )

        # 4. SHEET SHEET3
        ws_sum = wb.create_sheet(title="Sheet3")
        for c_idx, col_name in enumerate(df_sheet3.columns, 1):
            cell = ws_sum.cell(row=1, column=c_idx, value=col_name)
            cell.fill = grey_fill

        for s_idx, r in enumerate(df_sheet3.itertuples(index=False), 2):
            ws_sum.cell(row=s_idx, column=1, value=r[0])
            ws_sum.cell(row=s_idx, column=2, value=r[1])
            ws_sum.cell(row=s_idx, column=3, value=r[2])

        last_s3_row = len(df_sheet3) + 1
        s3_gt_row = last_s3_row + 1

        s3_gt1 = ws_sum.cell(row=s3_gt_row, column=1, value="Grand Total")
        s3_gt2 = ws_sum.cell(
            row=s3_gt_row, column=2, value=f"=SUM(B2:B{last_s3_row})"
        )
        s3_gt3 = ws_sum.cell(
            row=s3_gt_row, column=3, value=f"=SUM(C2:C{last_s3_row})"
        )

        for cell in (s3_gt1, s3_gt2, s3_gt3):
            cell.font = Font(bold=True, name="Calibri")
            cell.fill = grey_fill

        apply_table_formatting(
            ws_sum, start_row=1, max_col=len(df_sheet3.columns)
        )

        # Export Excel
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        st.download_button(
            label="📥 Download File Excel SJM & Marking",
            data=output,
            file_name=f"PERFECT_SJ_MANUAL_{uploaded_file.name.replace('.pdf', '')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
